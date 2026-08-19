"""Word .docx extraction.

Uses paragraph styles for structure: Word's built-in Heading 1/2/3 styles become
chapter boundaries. Documents that never use heading styles come back as one chapter.

Note: .doc (the pre-2007 binary format) is NOT supported — convert it first with
`libreoffice --headless --convert-to docx file.doc`.
"""
from __future__ import annotations

from pathlib import Path

from ..document import Chapter, Document


def extract_docx(path: Path) -> Document:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise ImportError(
            "Word support needs `python-docx` (pip install 'narratortool[docx]')"
        ) from exc

    f = docx.Document(str(path))

    props = f.core_properties
    doc = Document(
        title=(props.title or "").strip() or None,
        author=(props.author or "").strip() or None,
    )

    chapters: list[Chapter] = []
    current_title: str | None = None
    buf: list[str] = []

    for para in f.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        style = (para.style.name or "") if para.style is not None else ""
        if style.startswith("Heading") and _heading_level(style) <= 3:
            if buf:
                chapters.append(Chapter(text="\n".join(buf), title=current_title))
                buf = []
            current_title = text
            continue

        buf.append(text)

    if buf:
        chapters.append(Chapter(text="\n".join(buf), title=current_title))

    doc.chapters = chapters
    return doc


def _heading_level(style_name: str) -> int:
    tail = style_name.replace("Heading", "").strip()
    return int(tail) if tail.isdigit() else 99
